import csv
import hashlib
import io
from datetime import datetime

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db.models import Count, Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .forms import ImportCommentsForm
from .models import Classification, Comment, Profile, Publication
from .services.qualification import analyze_comment


def _lead_queryset():
    return (
        Classification.objects.filter(is_lead=True)
        .select_related('comment__publication__profile')
        .order_by('-comment__captured_at')
    )


@login_required
def dashboard(request):
    leads = _lead_queryset()
    context = {
        'profiles': Profile.objects.filter(active=True).count(),
        'publications': Publication.objects.count(),
        'comments': Comment.objects.count(),
        'leads': leads.count(),
        'high': leads.filter(qualification='high').count(),
        'medium': leads.filter(qualification='medium').count(),
        'low': leads.filter(qualification='low').count(),
        'pending': leads.filter(review_status='pending').count(),
        'recent_leads': leads[:8],
        'profile_summary': (
            Profile.objects.annotate(
                publication_count=Count('publications', distinct=True),
                comment_count=Count('publications__comments', distinct=True),
                lead_count=Count(
                    'publications__comments__classification',
                    filter=Q(publications__comments__classification__is_lead=True),
                    distinct=True,
                ),
            )
            .order_by('-lead_count', 'handle')[:6]
        ),
    }
    return render(request, 'dashboard.html', context)


@login_required
def leads_list(request):
    queryset = _lead_queryset()
    level = request.GET.get('nivel', '').strip()
    search = request.GET.get('q', '').strip()
    profile = request.GET.get('perfil', '').strip()

    if level in {'high', 'medium', 'low'}:
        queryset = queryset.filter(qualification=level)
    if profile:
        queryset = queryset.filter(comment__publication__profile__handle=profile)
    if search:
        queryset = queryset.filter(
            Q(comment__username__icontains=search)
            | Q(comment__display_name__icontains=search)
            | Q(comment__text_original__icontains=search)
            | Q(comment__publication__property_type__icontains=search)
            | Q(comment__publication__property_region__icontains=search)
        )

    context = {
        'items': queryset,
        'profiles_filter': Profile.objects.order_by('handle'),
        'selected_level': level,
        'selected_profile': profile,
        'search': search,
        'total': queryset.count(),
    }
    return render(request, 'leads.html', context)


@login_required
def profiles_list(request):
    profiles = Profile.objects.annotate(
        publication_count=Count('publications', distinct=True),
        comment_count=Count('publications__comments', distinct=True),
        lead_count=Count(
            'publications__comments__classification',
            filter=Q(publications__comments__classification__is_lead=True),
            distinct=True,
        ),
    ).order_by('handle')
    return render(request, 'profiles.html', {'profiles_list': profiles})


@login_required
def review_list(request):
    items = _lead_queryset().filter(review_status='pending')
    return render(request, 'review.html', {'items': items, 'total': items.count()})


@login_required
def review_action(request, classification_id, action):
    if request.method != 'POST':
        return redirect('review')
    item = get_object_or_404(Classification, pk=classification_id)
    if action == 'approve':
        item.review_status = 'approved'
        item.is_lead = True
        messages.success(request, 'Lead aprovado.')
    elif action == 'reject':
        item.review_status = 'rejected'
        item.is_lead = False
        messages.info(request, 'Registro rejeitado como lead.')
    item.save(update_fields=['review_status', 'is_lead', 'updated_at'])
    return redirect('review')


@login_required
def review_bulk_approve(request):
    if request.method != 'POST':
        return redirect('review')

    classification_ids = request.POST.getlist('classification_ids')
    if not classification_ids:
        messages.warning(request, 'Selecione pelo menos um lead para aprovar.')
        return redirect('review')

    queryset = Classification.objects.filter(
        pk__in=classification_ids,
        review_status='pending',
        is_lead=True,
    )
    count = queryset.update(
        review_status='approved',
        is_lead=True,
        updated_at=timezone.now(),
    )

    if count:
        messages.success(request, f'{count} lead(s) aprovado(s) de uma vez.')
    else:
        messages.info(request, 'Nenhum lead pendente foi alterado.')
    return redirect('review')


@login_required
def import_comments(request):
    form = ImportCommentsForm(request.POST or None, request.FILES or None)
    summary = None

    if request.method == 'POST' and form.is_valid():
        uploaded = form.cleaned_data['arquivo']
        raw = uploaded.read()
        try:
            text = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = raw.decode('latin-1')

        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=';,')
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = ';'

        reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
        new_count = duplicate_count = invalid_count = 0

        for row in reader:
            try:
                origin = (row.get('perfil_origem') or '').strip()
                handle = (row.get('handle_usuario') or '').strip()
                comment_text = (row.get('comentario') or '').strip()
                permalink = (row.get('link_publicacao') or '').strip()
                date_display = (row.get('data_comentario') or '').strip()
                if not origin or not handle or not comment_text or not permalink:
                    invalid_count += 1
                    continue

                if not origin.startswith('@'):
                    origin = '@' + origin
                if not handle.startswith('@'):
                    handle = '@' + handle

                profile_obj, _ = Profile.objects.get_or_create(
                    handle=origin,
                    defaults={'source_type': 'manual'},
                )
                publication, _ = Publication.objects.get_or_create(
                    profile=profile_obj,
                    permalink=permalink,
                    defaults={
                        'property_type': (row.get('tipo_imovel') or '').strip(),
                        'property_region': (row.get('regiao') or '').strip(),
                        'transaction_type': (row.get('finalidade') or '').strip(),
                        'caption_original': (row.get('descricao_publicacao') or '').strip(),
                    },
                )

                raw_key = '|'.join([origin, permalink, handle, comment_text, date_display])
                fingerprint = hashlib.sha256(raw_key.encode('utf-8')).hexdigest()
                if Comment.objects.filter(fingerprint=fingerprint).exists():
                    duplicate_count += 1
                    continue

                comment = Comment.objects.create(
                    publication=publication,
                    username=handle,
                    display_name=(row.get('nome_usuario') or '').strip(),
                    text_original=comment_text,
                    displayed_date_original=date_display,
                    source_type='manual',
                    fingerprint=fingerprint,
                )
                result = analyze_comment(comment_text)
                Classification.objects.create(
                    comment=comment,
                    is_lead=result['is_lead'],
                    intents=result.get('intents', []),
                    qualification=result['qualification'],
                    evidence=result.get('evidence', []),
                    confidence=result.get('confidence'),
                    exclusion_reason=result.get('exclusion_reason', ''),
                    review_status='pending' if result['is_lead'] else 'auto',
                    classifier_version='0.2.0',
                )
                new_count += 1
            except Exception:
                invalid_count += 1

        summary = {'new': new_count, 'duplicates': duplicate_count, 'invalid': invalid_count}
        messages.success(request, 'Importacao concluida.')
        form = ImportCommentsForm()

    return render(request, 'import.html', {'form': form, 'summary': summary})


@login_required
def export_excel(request):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Leads'

    headers = [
        'Nome do Usuario', '@Instagram', 'Link do Perfil', 'Comentario Original',
        'Link da Publicacao', 'Bairro/Tipo de Imovel', 'Sinal de Intencao Comercial',
        'Nivel de Qualificacao', 'Perfil de Origem', 'Data do Comentario',
    ]
    sheet.append(headers)
    header_fill = PatternFill('solid', fgColor='16324F')
    for cell in sheet[1]:
        cell.font = Font(color='FFFFFF', bold=True)
        cell.fill = header_fill

    for item in _lead_queryset():
        comment = item.comment
        pub = comment.publication
        property_context = ' - '.join(filter(None, [pub.property_type, pub.property_region]))
        sheet.append([
            comment.display_name or comment.username,
            comment.username,
            f"https://www.instagram.com/{comment.username.lstrip('@')}/",
            comment.text_original,
            pub.permalink,
            property_context,
            ', '.join(item.intents),
            item.get_qualification_display(),
            pub.profile.handle,
            comment.displayed_date_original or (comment.commented_at.isoformat() if comment.commented_at else ''),
        ])

    for column in sheet.columns:
        max_length = min(max(len(str(cell.value or '')) for cell in column) + 2, 55)
        sheet.column_dimensions[column[0].column_letter].width = max_length

    summary = workbook.create_sheet('Resumo')
    summary.append(['Gerado em', timezone.localtime().strftime('%d/%m/%Y %H:%M')])
    summary.append(['Perfis ativos', Profile.objects.filter(active=True).count()])
    summary.append(['Comentarios', Comment.objects.count()])
    summary.append(['Leads', Classification.objects.filter(is_lead=True).count()])
    summary.append(['Alto', Classification.objects.filter(is_lead=True, qualification='high').count()])
    summary.append(['Medio', Classification.objects.filter(is_lead=True, qualification='medium').count()])
    summary.append(['Baixo', Classification.objects.filter(is_lead=True, qualification='low').count()])

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="ChaveRadar_leads.xlsx"'
    workbook.save(response)
    return response
